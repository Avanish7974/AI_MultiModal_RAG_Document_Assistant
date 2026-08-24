from docx import Document


def extract_docx_text(file_path):
    """
    Extract text from DOCX paragraphs and tables.

    Returns:
        list of dictionaries containing:
        - text
        - location
    """

    document = Document(file_path)

    extracted_data = []

    # =====================================================
    # PARAGRAPHS
    # =====================================================

    for index, paragraph in enumerate(
        document.paragraphs,
        start=1
    ):

        text = paragraph.text.strip()

        if text:

            extracted_data.append(
                {
                    "text": text,
                    "location": index
                }
            )

    # =====================================================
    # TABLES
    # =====================================================

    for table_index, table in enumerate(
        document.tables,
        start=1
    ):

        for row_index, row in enumerate(
            table.rows,
            start=1
        ):

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:

                extracted_data.append(
                    {
                        "text": " | ".join(row_text),
                        "location": f"Table {table_index}, Row {row_index}"
                    }
                )

    return extracted_data