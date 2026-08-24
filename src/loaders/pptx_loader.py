from pptx import Presentation


def extract_pptx_text(file_path):
    """
    Extract text from PPTX slides.

    Returns:
        list of dictionaries containing:
        - text
        - location (slide number)
    """

    presentation = Presentation(file_path)

    extracted_data = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            # -------------------------------------------------
            # NORMAL TEXT
            # -------------------------------------------------

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_text.append(text)

            # -------------------------------------------------
            # TABLE
            # -------------------------------------------------

            if shape.has_table:

                table = shape.table

                for row in table.rows:

                    row_text = []

                    for cell in row.cells:

                        cell_text = cell.text.strip()

                        if cell_text:
                            row_text.append(cell_text)

                    if row_text:

                        slide_text.append(
                            " | ".join(row_text)
                        )

        combined_text = "\n".join(slide_text).strip()

        if combined_text:

            extracted_data.append(
                {
                    "text": combined_text,
                    "location": slide_number
                }
            )

    return extracted_data